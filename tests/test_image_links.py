from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

import image_links


def _configure_tmp_project(monkeypatch, tmp_path: Path) -> Path:
    outputs = tmp_path / "outputs"
    (outputs / "csv").mkdir(parents=True)
    (outputs / "excel").mkdir(parents=True)
    (outputs / "images").mkdir(parents=True)

    monkeypatch.setattr(image_links, "ROOT", tmp_path)
    monkeypatch.setattr(image_links, "CONFIG_PATH", tmp_path / "config.yaml")
    monkeypatch.setattr(image_links, "REQUEST_PATH", tmp_path / "ADAPTER_REQUEST.md")
    monkeypatch.setattr(image_links, "OUTPUTS_DIR", outputs)
    monkeypatch.setattr(image_links, "IMAGES_DIR", outputs / "images")

    (tmp_path / "config.yaml").write_text("image_public_base_url: ''\n", encoding="utf-8")
    return outputs


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    fieldnames = list(rows[0])
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def test_moves_misplaced_brand_folder_and_writes_template_links(monkeypatch, tmp_path):
    outputs = _configure_tmp_project(monkeypatch, tmp_path)
    (tmp_path / "ADAPTER_REQUEST.md").write_text(
        "Domena zdjęć: "
        "https://mojadomena.pl/produkty/[marka albo kategoria]/[sku].[rozszerzenie]\n",
        encoding="utf-8",
    )
    _write_csv(
        outputs / "csv" / "products.csv",
        [{"sku": "ABC123", "brand": "Test Brand", "category": "Ignored"}],
    )

    misplaced = outputs / "Test Brand"
    misplaced.mkdir()
    (misplaced / "ABC123.jpg").write_bytes(b"first")
    (misplaced / "ABC123_1.webp").write_bytes(b"second")

    assert image_links.add_public_image_links() == 1

    assert not misplaced.exists()
    assert (outputs / "images" / "test-brand" / "ABC123.jpg").exists()
    assert (outputs / "images" / "test-brand" / "ABC123_1.webp").exists()

    row = _read_csv(outputs / "csv" / "products.csv")[0]
    assert row["zdj1"] == "https://mojadomena.pl/produkty/test-brand/ABC123.jpg"
    assert row["zdj2"] == "https://mojadomena.pl/produkty/test-brand/ABC123_1.webp"


def test_moves_loose_images_under_category_and_writes_base_url_links(monkeypatch, tmp_path):
    outputs = _configure_tmp_project(monkeypatch, tmp_path)
    (tmp_path / "ADAPTER_REQUEST.md").write_text(
        "Domena zdjęć: https://cdn.example.pl/produkty\n",
        encoding="utf-8",
    )
    _write_csv(
        outputs / "csv" / "products.csv",
        [{"SKU": "SKU-9", "Marka": "", "Kategoria": "Łóżka dziecięce"}],
    )
    (outputs / "images" / "SKU-9.png").write_bytes(b"image")

    assert image_links.add_public_image_links() == 1

    assert not (outputs / "images" / "SKU-9.png").exists()
    assert (outputs / "images" / "lozka-dzieciece" / "SKU-9.png").exists()

    row = _read_csv(outputs / "csv" / "products.csv")[0]
    assert row["zdj1"] == "https://cdn.example.pl/produkty/lozka-dzieciece/SKU-9.png"


def test_updates_excel_with_configured_segment(monkeypatch, tmp_path):
    outputs = _configure_tmp_project(monkeypatch, tmp_path)
    (tmp_path / "ADAPTER_REQUEST.md").write_text(
        "Domena zdjęć "
        "https://mojadomena.pl/produkty/[marka albo kategoria]/[sku].[rozszerzenie]\n"
        "Marka: Stała Marka\n",
        encoding="utf-8",
    )

    workbook_path = outputs / "excel" / "products.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["sku", "brand", "category"])
    ws.append(["XYZ", "Other Brand", "Other Category"])
    wb.save(workbook_path)

    (outputs / "images" / "XYZ.jpg").write_bytes(b"image")

    assert image_links.add_public_image_links() == 1

    assert (outputs / "images" / "stala-marka" / "XYZ.jpg").exists()

    wb = openpyxl.load_workbook(workbook_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers[-1] == "zdj1"
    link_cell = ws.cell(row=2, column=headers.index("zdj1") + 1)
    assert link_cell.value == "https://mojadomena.pl/produkty/stala-marka/XYZ.jpg"
    assert link_cell.hyperlink.target == link_cell.value
    wb.close()
