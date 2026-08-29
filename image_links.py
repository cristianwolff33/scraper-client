from __future__ import annotations

import csv
import filecmp
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote, unquote, urlsplit

import yaml


ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.yaml"
REQUEST_PATH = ROOT / "ADAPTER_REQUEST.md"
OUTPUTS_DIR = ROOT / "outputs"
IMAGES_DIR = OUTPUTS_DIR / "images"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
RESERVED_OUTPUT_DIRS = {"cache", "csv", "excel", "images", "json", "logs", "xml"}
DEFAULT_IMAGE_SEGMENT = "bez-kategorii"
ASCII_TRANSLATION = str.maketrans(
    {
        "ą": "a",
        "ć": "c",
        "ę": "e",
        "ł": "l",
        "ń": "n",
        "ó": "o",
        "ś": "s",
        "ź": "z",
        "ż": "z",
        "Ą": "A",
        "Ć": "C",
        "Ę": "E",
        "Ł": "L",
        "Ń": "N",
        "Ó": "O",
        "Ś": "S",
        "Ź": "Z",
        "Ż": "Z",
    }
)

SKU_FIELDS = {
    "sku",
    "kod",
    "kod producenta",
    "kod producenta produktu",
    "kod_producenta",
    "product code",
    "producer code",
}
BRAND_FIELDS = {"brand", "marka", "producent", "manufacturer"}
CATEGORY_FIELDS = {"category", "kategoria", "categories", "breadcrumb"}
PRODUCTS_PATH_SEGMENTS = {"produkty", "products"}


@dataclass(frozen=True)
class ImageLinkSettings:
    public_url: str
    configured_segment: str

    @property
    def uses_template(self) -> bool:
        return bool(re.search(r"\[[^\]]+\]", self.public_url))


def _ascii_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").translate(ASCII_TRANSLATION))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return text


def _normalize_label(value: object) -> str:
    text = _ascii_text(value)
    text = text.lower().replace("_", " ")
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _safe_sku(value: object) -> str:
    text = str(value or "").strip()
    return re.sub(r"[^\w\-]", "_", text).strip("_") or "unnamed"


def _safe_segment(value: object) -> str:
    text = _ascii_text(value)
    text = re.sub(r"[^A-Za-z0-9]+", "-", text.lower()).strip("-")
    return text or DEFAULT_IMAGE_SEGMENT


def _read_config() -> dict:
    if not CONFIG_PATH.exists():
        return {}
    with CONFIG_PATH.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def _extract_url(value: str) -> str:
    cleaned = value.strip().strip("\"'")
    if cleaned.startswith(("http://", "https://")):
        return cleaned.rstrip(".,")

    match = re.search(r"https?://", value)
    if match:
        return value[match.start() :].strip().rstrip(".,")
    return cleaned


def _read_adapter_request() -> dict[str, str]:
    if not REQUEST_PATH.exists():
        return {}

    aliases = {
        "domena": "image_public_base_url",
        "domena zdjec": "image_public_base_url",
        "domena obrazow": "image_public_base_url",
        "wzor domeny": "image_public_base_url",
        "wzor domeny zdjec": "image_public_base_url",
        "public image domain": "image_public_base_url",
        "image domain": "image_public_base_url",
        "marka": "image_public_segment",
        "brand": "image_public_segment",
        "segment": "image_public_segment",
        "kategoria": "image_public_segment",
        "category": "image_public_segment",
    }

    values: dict[str, str] = {}
    for raw_line in REQUEST_PATH.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip().lstrip("-").strip()
        if not line:
            continue

        url_match = re.search(r"https?://", line)
        separator_index = line.find(":")
        if separator_index != -1 and (
            url_match is None or separator_index < url_match.start()
        ):
            key = line[:separator_index]
            value = line[separator_index + 1 :]
            mapped = aliases.get(_normalize_label(key))
            if mapped and value.strip():
                values[mapped] = _extract_url(value) if mapped.endswith("_url") else value.strip()
            continue

        if url_match:
            key = line[: url_match.start()]
            value = line[url_match.start() :]
            mapped = aliases.get(_normalize_label(key))
            if mapped:
                values[mapped] = _extract_url(value) if mapped.endswith("_url") else value
            continue

        # Accept short client prompts like: "Domena https://example.pl/produkty/..."
        match = re.match(r"^(domena|marka|brand|kategoria|category)\s+(.+)$", line, re.IGNORECASE)
        if match:
            mapped = aliases.get(_normalize_label(match.group(1)))
            value = match.group(2).strip()
            if mapped and value:
                values[mapped] = _extract_url(value) if mapped.endswith("_url") else value

    return values


def _settings() -> ImageLinkSettings:
    config = _read_config()
    request = _read_adapter_request()
    public_url = (
        request.get("image_public_base_url")
        or config.get("image_public_base_url")
        or config.get("image_public_domain")
        or ""
    )
    configured_segment = (
        request.get("image_public_segment")
        or request.get("image_public_brand")
        or config.get("image_public_segment")
        or config.get("image_public_brand")
        or config.get("brand")
        or ""
    )
    return ImageLinkSettings(
        public_url=str(public_url).strip().rstrip("/"),
        configured_segment=_safe_segment(configured_segment) if configured_segment else "",
    )


def _image_base_sku(path: Path) -> str:
    return re.sub(r"_\d+$", "", path.stem)


def _sku_lookup_keys(sku: object) -> list[str]:
    raw = str(sku or "").strip()
    keys = [raw, _safe_sku(raw), _image_base_sku(Path(_safe_sku(raw)))]
    return [key for index, key in enumerate(keys) if key and key not in keys[:index]]


def _find_field(fieldnames: list[object], aliases: set[str]) -> object | None:
    normalized = {_normalize_label(name): name for name in fieldnames if name is not None}
    for alias in aliases:
        found = normalized.get(_normalize_label(alias))
        if found is not None:
            return found
    return None


def _row_value(row: dict[object, object], field: object | None) -> str:
    if field is None:
        return ""
    value = row.get(field)
    return str(value or "").strip()


def _segment_from_values(brand: str, category: str, settings: ImageLinkSettings) -> str:
    if settings.configured_segment:
        return settings.configured_segment
    return _safe_segment(brand or category)


def _collect_segments_from_csv(path: Path, settings: ImageLinkSettings) -> dict[str, str]:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = list(reader.fieldnames or [])

    sku_field = _find_field(fieldnames, SKU_FIELDS)
    if not sku_field:
        return {}

    brand_field = _find_field(fieldnames, BRAND_FIELDS)
    category_field = _find_field(fieldnames, CATEGORY_FIELDS)
    segments: dict[str, str] = {}
    for row in rows:
        sku = _row_value(row, sku_field)
        if not sku:
            continue
        segment = _segment_from_values(
            _row_value(row, brand_field),
            _row_value(row, category_field),
            settings,
        )
        for key in _sku_lookup_keys(sku):
            segments.setdefault(key, segment)
    return segments


def _collect_segments_from_excel(path: Path, settings: ImageLinkSettings) -> dict[str, str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    sku_field = _find_field(headers, SKU_FIELDS)
    if not sku_field:
        wb.close()
        return {}

    brand_field = _find_field(headers, BRAND_FIELDS)
    category_field = _find_field(headers, CATEGORY_FIELDS)
    sku_idx = headers.index(sku_field)
    brand_idx = (
        headers.index(brand_field)
        if brand_field is not None and brand_field in headers
        else None
    )
    category_idx = (
        headers.index(category_field)
        if category_field is not None and category_field in headers
        else None
    )

    segments: dict[str, str] = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        sku = str(values[sku_idx] or "").strip()
        if not sku:
            continue
        brand = str(values[brand_idx] or "").strip() if brand_idx is not None else ""
        category = str(values[category_idx] or "").strip() if category_idx is not None else ""
        segment = _segment_from_values(brand, category, settings)
        for key in _sku_lookup_keys(sku):
            segments.setdefault(key, segment)
    wb.close()
    return segments


def _collect_product_segments(settings: ImageLinkSettings) -> dict[str, str]:
    segments: dict[str, str] = {}
    for path in sorted((OUTPUTS_DIR / "csv").rglob("*.csv")):
        segments.update(_collect_segments_from_csv(path, settings))
    for path in sorted((OUTPUTS_DIR / "excel").rglob("*.xlsx")):
        segments.update(_collect_segments_from_excel(path, settings))
    return segments


def _next_image_variant(target_dir: Path, image_path: Path) -> Path:
    base_sku = _image_base_sku(image_path)
    suffix = image_path.suffix
    index = 0
    while True:
        stem = base_sku if index == 0 else f"{base_sku}_{index}"
        candidate = target_dir / f"{stem}{suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def _move_image(path: Path, target_dir: Path) -> bool:
    target_dir.mkdir(parents=True, exist_ok=True)
    if path.parent == target_dir:
        return False

    target = target_dir / path.name
    if target.exists():
        if filecmp.cmp(path, target, shallow=False):
            path.unlink()
            return True
        target = _next_image_variant(target_dir, path)

    shutil.move(str(path), str(target))
    return True


def _remove_empty_dirs(path: Path) -> None:
    if not path.exists():
        return
    for child in sorted(path.rglob("*"), key=lambda item: len(item.parts), reverse=True):
        if child.is_dir():
            try:
                child.rmdir()
            except OSError:
                pass


def _move_misplaced_output_dirs() -> int:
    if not OUTPUTS_DIR.exists():
        return 0

    moved = 0
    IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    for child in OUTPUTS_DIR.iterdir():
        if not child.is_dir() or child.name.lower() in RESERVED_OUTPUT_DIRS:
            continue

        target_dir = IMAGES_DIR / _safe_segment(child.name)
        for path in sorted(child.rglob("*")):
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
                moved += int(_move_image(path, target_dir))
        _remove_empty_dirs(child)
        try:
            child.rmdir()
        except OSError:
            pass
    return moved


def _move_loose_images(product_segments: dict[str, str], settings: ImageLinkSettings) -> int:
    if not IMAGES_DIR.exists():
        return 0

    moved = 0
    for path in sorted(IMAGES_DIR.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        base_sku = _image_base_sku(path)
        row_segment = settings.configured_segment
        for key in _sku_lookup_keys(base_sku):
            row_segment = row_segment or product_segments.get(key, "")
        row_segment = row_segment or DEFAULT_IMAGE_SEGMENT
        segment = _link_segment_for(path, settings, row_segment)
        moved += int(_move_image(path, IMAGES_DIR / segment))

    _remove_empty_dirs(IMAGES_DIR)
    return moved


def _organize_images(product_segments: dict[str, str], settings: ImageLinkSettings) -> int:
    moved = _move_misplaced_output_dirs()
    moved += _move_loose_images(product_segments, settings)
    return moved


def _image_index() -> dict[str, list[Path]]:
    index: dict[str, list[Path]] = {}
    if not IMAGES_DIR.exists():
        return index

    for path in sorted(IMAGES_DIR.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        base_sku = _image_base_sku(path)
        for key in _sku_lookup_keys(base_sku):
            index.setdefault(key, []).append(path)

    def sort_key(path: Path) -> tuple[str, int, str]:
        match = re.search(r"_(\d+)$", path.stem)
        order = int(match.group(1)) if match else 0
        segment = path.parent.name if path.parent != IMAGES_DIR else ""
        return segment, order, path.name

    for paths in index.values():
        paths.sort(key=sort_key)
    return index


def _template_url_for(path: Path, settings: ImageLinkSettings, segment: str) -> str:
    replacements = {
        "marka": segment,
        "brand": segment,
        "kategoria": segment,
        "category": segment,
        "marka albo kategoria": segment,
        "marka lub kategoria": segment,
        "brand or category": segment,
        "sku": path.stem,
        "rozszerzenie": path.suffix.lstrip("."),
        "extension": path.suffix.lstrip("."),
        "ext": path.suffix.lstrip("."),
    }

    def replace(match: re.Match[str]) -> str:
        key = _normalize_label(match.group(1))
        value = replacements.get(key)
        if value is None:
            return match.group(0)
        return quote(value, safe="")

    return re.sub(r"\[([^\]]+)\]", replace, settings.public_url)


def _segment_after_products(url: str) -> str:
    parts = [part for part in urlsplit(url).path.split("/") if part]
    for index, part in enumerate(parts[:-1]):
        if _normalize_label(unquote(part)) in PRODUCTS_PATH_SEGMENTS:
            next_part = unquote(parts[index + 1])
            if "[" not in next_part and "]" not in next_part:
                return _safe_segment(next_part)
    return ""


def _link_segment_for(path: Path, settings: ImageLinkSettings, row_segment: str) -> str:
    segment = settings.configured_segment or row_segment or DEFAULT_IMAGE_SEGMENT
    if settings.uses_template:
        rendered_url = _template_url_for(path, settings, segment)
        return _segment_after_products(rendered_url) or segment
    return _segment_after_products(settings.public_url) or segment


def _url_for(path: Path, settings: ImageLinkSettings, row_segment: str) -> str:
    if not settings.public_url:
        return ""

    segment = _link_segment_for(path, settings, row_segment)

    if settings.uses_template:
        return _template_url_for(path, settings, segment)

    file_segment = quote(path.name, safe="")
    fixed_segment = _segment_after_products(settings.public_url)
    if fixed_segment:
        return f"{settings.public_url}/{file_segment}"

    segment_segment = quote(segment, safe="")
    return f"{settings.public_url}/{segment_segment}/{file_segment}"


def _column_name(index: int) -> str:
    return f"zdj{index}"


def _images_for_sku(sku: str, images: dict[str, list[Path]]) -> list[Path]:
    seen: set[Path] = set()
    matched: list[Path] = []
    for key in _sku_lookup_keys(sku):
        for path in images.get(key, []):
            if path not in seen:
                seen.add(path)
                matched.append(path)
    return matched


def _links_for_sku(
    sku: str,
    row_brand: str,
    row_category: str,
    settings: ImageLinkSettings,
    images: dict[str, list[Path]],
) -> list[str]:
    row_segment = _segment_from_values(row_brand, row_category, settings)
    return [_url_for(path, settings, row_segment) for path in _images_for_sku(sku, images)]


def _update_csv(path: Path, settings: ImageLinkSettings, images: dict[str, list[Path]]) -> bool:
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = list(reader.fieldnames or [])

    sku_field = _find_field(fieldnames, SKU_FIELDS)
    if not rows or not sku_field:
        return False

    brand_field = _find_field(fieldnames, BRAND_FIELDS)
    category_field = _find_field(fieldnames, CATEGORY_FIELDS)
    max_links = 0
    for row in rows:
        links = _links_for_sku(
            _row_value(row, sku_field),
            _row_value(row, brand_field),
            _row_value(row, category_field),
            settings,
            images,
        )
        max_links = max(max_links, len(links))
        for index, link in enumerate(links, start=1):
            row[_column_name(index)] = link

    for index in range(1, max_links + 1):
        name = _column_name(index)
        if name not in fieldnames:
            fieldnames.append(name)

    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return max_links > 0


def _update_excel(path: Path, settings: ImageLinkSettings, images: dict[str, list[Path]]) -> bool:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    sku_field = _find_field(headers, SKU_FIELDS)
    if not sku_field:
        wb.close()
        return False

    brand_field = _find_field(headers, BRAND_FIELDS)
    category_field = _find_field(headers, CATEGORY_FIELDS)
    sku_col = headers.index(sku_field) + 1
    brand_col = (
        headers.index(brand_field) + 1
        if brand_field is not None and brand_field in headers
        else None
    )
    category_col = (
        headers.index(category_field) + 1
        if category_field is not None and category_field in headers
        else None
    )
    normalized_headers = [str(header).strip() if header is not None else "" for header in headers]
    max_links = 0

    for row_idx in range(2, ws.max_row + 1):
        sku = str(ws.cell(row=row_idx, column=sku_col).value or "")
        row_brand = str(ws.cell(row=row_idx, column=brand_col).value or "") if brand_col else ""
        row_category = (
            str(ws.cell(row=row_idx, column=category_col).value or "")
            if category_col
            else ""
        )
        links = _links_for_sku(sku, row_brand, row_category, settings, images)
        max_links = max(max_links, len(links))
        for index, link in enumerate(links, start=1):
            header = _column_name(index)
            if header not in normalized_headers:
                ws.cell(row=1, column=len(normalized_headers) + 1, value=header)
                normalized_headers.append(header)
            cell = ws.cell(row=row_idx, column=normalized_headers.index(header) + 1)
            cell.value = link
            cell.hyperlink = link
            cell.style = "Hyperlink"

    wb.save(path)
    wb.close()
    return max_links > 0


def add_public_image_links() -> int:
    settings = _settings()
    product_segments = _collect_product_segments(settings)
    moved = _organize_images(product_segments, settings)

    if moved:
        print(f"Image organization complete: moved {moved} image file(s).")

    if not settings.public_url:
        print(
            "Image link generation skipped: set Domena zdjęć in ADAPTER_REQUEST.md "
            "or image_public_base_url in config.yaml."
        )
        return 0

    images = _image_index()
    if not images:
        print("Image link generation skipped: no images found in outputs/images.")
        return 0

    updated = 0
    for path in sorted((OUTPUTS_DIR / "csv").rglob("*.csv")):
        updated += int(_update_csv(path, settings, images))
    for path in sorted((OUTPUTS_DIR / "excel").rglob("*.xlsx")):
        updated += int(_update_excel(path, settings, images))

    print(f"Image link generation complete: updated {updated} export file(s).")
    return updated


if __name__ == "__main__":
    add_public_image_links()
