from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import json

from core.exporter import BaseExporter
from models.product import Product

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


_HEADERS = [
    ("SKU", "sku"),
    ("Local Symbol", "local_symbol"),
    ("EAN", "ean"),
    ("Name", "name"),
    ("Brand", "brand"),
    ("Category", "category"),
    ("Description", "description"),
    ("Price", "price"),
    ("Currency", "currency"),
    ("Stock", "stock"),
    ("Unit", "unit"),
    ("Availability", "availability"),
    ("Images", "image_urls"),
    ("zdj 1", "image_url_1"),
    ("Match Status", "match_status"),
    ("Match Method", "match_method"),
    ("Confidence", "match_confidence"),
    ("Source Name", "source_name"),
    ("Description Source", "description_source"),
    ("Image Source", "image_source"),
    ("External Sources", "external_sources"),
    ("Not Found Reason", "reason_if_not_found"),
    ("Source URL", "source_url"),
]


class ExcelExporter(BaseExporter):
    """Exports products to .xlsx with header styling."""

    def export(self, products: list[Product], destination: Path) -> Path:
        if not HAS_OPENPYXL:
            raise RuntimeError("openpyxl is required: pip install openpyxl")

        destination.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = destination / f"products_{ts}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        rows = [_flatten_product(product) for product in products]
        headers = _headers(rows)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2E5090")
        header_align = Alignment(horizontal="center")
        wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        description_col = next(
            (i for i, (_, field) in enumerate(headers, start=1) if field == "description"),
            None,
        )

        for col_idx, (header, _) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for row_idx, d in enumerate(rows, start=2):
            for col_idx, (_, field) in enumerate(headers, start=1):
                value = d.get(field, "")
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if col_idx == description_col:
                    cell.alignment = wrap_align
            if description_col:
                line_count = str(d.get("description", "")).count("\n") + 1
                ws.row_dimensions[row_idx].height = min(15 * line_count, 300)

        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].auto_size = True
        if description_col:
            ws.column_dimensions[get_column_letter(description_col)].width = 70

        wb.save(out)
        return out


def _headers(rows: list[dict[str, Any]]) -> list[tuple[str, str]]:
    fixed_fields = {field for _, field in _HEADERS}
    extra_fields: set[str] = set()
    for row in rows:
        extra_fields.update(row)
    return _HEADERS + [
        (field, field) for field in sorted(extra_fields.difference(fixed_fields))
    ]


def _flatten_product(product: Product) -> dict[str, Any]:
    row = product.to_dict()
    parameters = row.pop("parameters", {}) or {}
    attributes = row.pop("attributes", {}) or {}
    metadata = row.pop("metadata", {}) or {}
    export_image_urls = _export_image_urls(product, metadata)

    row["image_urls"] = "|".join(export_image_urls)
    row["downloaded_images"] = "|".join(product.downloaded_images)
    row["image_url_1"] = export_image_urls[0] if export_image_urls else ""
    row["parameters_json"] = _json(parameters)
    row["attributes_json"] = _json(attributes)
    row["metadata_json"] = _json(metadata)

    if export_image_urls:
        for index, url in enumerate(export_image_urls, start=1):
            row.setdefault(f"zdj{index}", url)

    for source in (parameters, metadata):
        if not isinstance(source, dict):
            continue
        for key, value in source.items():
            clean_key = str(key).strip()
            if clean_key and clean_key not in row:
                row[clean_key] = _serialize(value)

    return {key: _serialize(value) for key, value in row.items()}


def _export_image_urls(product: Product, metadata: dict[str, Any]) -> list[str]:
    public_urls = metadata.get("public_image_urls")
    if isinstance(public_urls, list):
        urls = [str(url) for url in public_urls if url]
        if urls:
            return urls
    return product.image_urls


def _serialize(value: Any) -> Any:
    if isinstance(value, list):
        return "|".join(str(item) for item in value)
    if isinstance(value, dict):
        return _json(value)
    return value


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str) if value else ""
